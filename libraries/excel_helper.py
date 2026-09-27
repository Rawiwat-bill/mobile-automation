from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from datetime import datetime
from zoneinfo import ZoneInfo
import json
import re

MAX_EXCEL_CELL = 30000


class ExcelHelper:
    ROBOT_LIBRARY_SCOPE = "GLOBAL"

    def __init__(self):
        self.wb = None
        self.ws = None
        self.path = None
        self.result_start_col = 7
        self.run_number = 1

    def load_workbook(self, path, sheet_name="APITest"):
        self.path = path
        self.wb = load_workbook(path)

        if sheet_name not in self.wb.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found. Available sheets: {self.wb.sheetnames}"
            )

        self.ws = self.wb[sheet_name]
        self._prepare_result_columns()

    def _prepare_result_columns(self):
        last_col = 0
        max_run = 0
        max_run_col = 0

        for cell in self.ws[1]:
            if cell.value is None:
                continue
            value = str(cell.value).strip()
            if cell.column > last_col:
                last_col = cell.column
            match = re.match(r"^ActualStatus_?(\d*)$", value)
            if match:
                run = int(match.group(1)) if match.group(1) else 1
                if run > max_run:
                    max_run = run
                    max_run_col = cell.column
                if not match.group(1):
                    for offset, name in enumerate(["ActualStatus", "ActualBody", "Result", "TestDate"]):
                        self._write_header(cell.column + offset, f"{name}_{run}")

        if max_run > 0:
            if self._block_has_data(max_run_col):
                self.run_number = max_run + 1
                self.result_start_col = last_col + 1
            else:
                self.run_number = max_run
                self.result_start_col = max_run_col
        else:
            self.run_number = 1
            self.result_start_col = max(7, last_col + 1)

        self._write_result_headers()

    def _block_has_data(self, start_col):
        for row in range(2, self.ws.max_row + 1):
            for offset in range(4):
                value = self.ws.cell(row=row, column=start_col + offset).value
                if value is not None and str(value).strip() != "":
                    return True
        return False

    def _write_result_headers(self):
        headers = ["ActualStatus", "ActualBody", "Result", "TestDate"]
        for offset, name in enumerate(headers):
            self._write_header(self.result_start_col + offset, f"{name}_{self.run_number}")

    def _write_header(self, column, value):
        cell = self.ws.cell(row=1, column=column, value=value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="top")

    def get_test_rows(self):
        if self.ws is None:
            raise RuntimeError("Worksheet is not loaded. Call Load Workbook first.")

        headers = []
        for cell in self.ws[1]:
            header = "" if cell.value is None else str(cell.value).strip()
            headers.append(header)

        input_columns = [
            i
            for i, header in enumerate(headers)
            if header
            and not re.match(r"^(ActualStatus|ActualBody|Result|TestDate)(_\d+)?$", header)
        ]

        rows = []
        for row_num, row in enumerate(self.ws.iter_rows(min_row=2, values_only=True), start=2):
            input_values = [row[i] if i < len(row) else None for i in input_columns]
            if all(v is None or str(v).strip() == "" for v in input_values):
                continue

            item = {"_row": row_num}
            for i, header in enumerate(headers):
                if not header:
                    continue
                value = row[i] if i < len(row) else None
                if isinstance(value, str):
                    value = value.strip()
                item[header] = value

            rows.append(item)

        return rows

    def compare_body(self, expected_body, actual_body):
        expected = "" if expected_body is None else str(expected_body).strip()
        actual = "" if actual_body is None else str(actual_body).strip()

        if expected == "":
            return True, "No expected body provided"

        try:
            exp_json = json.loads(expected)
            act_json = json.loads(actual)
        except Exception:
            return expected in actual, "Substring comparison"

        ok, detail = self._contains_json(act_json, exp_json)
        return ok, "JSON contains: " + detail

    def _contains_json(self, actual, expected):
        if isinstance(expected, dict):
            if not isinstance(actual, dict):
                return False, f"expected object but actual is {type(actual).__name__}"
            for key, exp_value in expected.items():
                if key not in actual:
                    return False, f"missing key '{key}'"
                ok, detail = self._contains_json(actual[key], exp_value)
                if not ok:
                    return False, f"key '{key}': {detail}"
            return True, "all expected keys found"
        if isinstance(expected, list):
            if not isinstance(actual, list):
                return False, f"expected array but actual is {type(actual).__name__}"
            for index, exp_item in enumerate(expected):
                if not any(self._contains_json(act_item, exp_item)[0] for act_item in actual):
                    return False, f"array item {index} not found"
            return True, "all expected array items found"
        return True, ""

    def get_test_date(self):
        return datetime.now(ZoneInfo("Asia/Bangkok")).strftime("%Y-%m-%d %H:%M:%S")

    def write_results(self, row_num, actual_status, actual_body, result, test_date):
        if self.ws is None:
            raise RuntimeError("Worksheet is not loaded. Call Load Workbook first.")

        start = self.result_start_col
        self.ws.cell(row=row_num, column=start, value=actual_status)
        self.ws.cell(row=row_num, column=start + 1, value=self._safe_text(actual_body))
        self.ws.cell(row=row_num, column=start + 2, value=result)
        self.ws.cell(row=row_num, column=start + 3, value=test_date)

    def save_workbook(self):
        if self.wb is None or self.path is None:
            return
        self.wb.save(self.path)

    def _safe_text(self, value):
        text = "" if value is None else str(value)
        if len(text) > MAX_EXCEL_CELL:
            return text[:MAX_EXCEL_CELL] + "...[truncated]"
        return text
